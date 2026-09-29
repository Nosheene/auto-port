"""Génère un CV Word en une colonne, lisible par un ATS."""

from docx import Document
from docx.enum.text import WD_LINE_SPACING, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.shared import Cm, Pt, RGBColor

ACCENT = "0E6B5C"
INK = "1C2430"
MUTED = "3E4C59"
RULE = "C5D5D1"
FONT = "Calibri"

OUTPUT = "/workspace/CV-Nosheene-Mohammad.docx"


def rgb(value: str) -> RGBColor:
    return RGBColor.from_string(value)


def set_run_font(run, size: float, bold: bool = False, color: str = INK, italic: bool = False) -> None:
    run.font.name = FONT
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = rgb(color)
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.find(qn("w:rFonts"))
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.append(r_fonts)
    for attribute in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        r_fonts.set(qn(attribute), FONT)
    lang = r_pr.find(qn("w:lang"))
    if lang is None:
        lang = OxmlElement("w:lang")
        r_pr.append(lang)
    lang.set(qn("w:val"), "fr-FR")
    lang.set(qn("w:eastAsia"), "fr-FR")


def set_paragraph_spacing(paragraph, before: float = 0, after: float = 0, line: int = 240) -> None:
    del line
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing_rule = WD_LINE_SPACING.SINGLE
    fmt.line_spacing = 1.0


def add_text(paragraph, text: str, size: float, bold: bool = False, color: str = INK, italic: bool = False):
    run = paragraph.add_run(text)
    set_run_font(run, size, bold, color, italic)
    return run


def add_hyperlink(paragraph, text: str, url: str, size: float) -> None:
    relation_id = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relation_id)
    run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")

    r_fonts = OxmlElement("w:rFonts")
    for attribute in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        r_fonts.set(qn(attribute), FONT)
    r_pr.append(r_fonts)

    size_el = OxmlElement("w:sz")
    size_el.set(qn("w:val"), str(int(size * 2)))
    r_pr.append(size_el)
    size_cs = OxmlElement("w:szCs")
    size_cs.set(qn("w:val"), str(int(size * 2)))
    r_pr.append(size_cs)

    color = OxmlElement("w:color")
    color.set(qn("w:val"), ACCENT)
    r_pr.append(color)

    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    underline.set(qn("w:color"), ACCENT)
    r_pr.append(underline)

    lang = OxmlElement("w:lang")
    lang.set(qn("w:val"), "fr-FR")
    r_pr.append(lang)

    run.append(r_pr)
    text_el = OxmlElement("w:t")
    text_el.set(qn("xml:space"), "preserve")
    text_el.text = text
    run.append(text_el)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_bottom_border(paragraph, color: str = RULE, size: str = "6") -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), color)
    borders.append(bottom)
    p_pr.append(borders)


def add_left_border(paragraph) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), "16")
    left.set(qn("w:space"), "10")
    left.set(qn("w:color"), ACCENT)
    borders.append(left)
    p_pr.append(borders)


def configure_document(document: Document) -> None:
    section = document.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.35)
    section.bottom_margin = Cm(1.25)
    section.left_margin = Cm(1.6)
    section.right_margin = Cm(1.6)
    section.header_distance = Cm(0.5)
    section.footer_distance = Cm(0.5)

    normal = document.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = rgb(INK)
    r_pr = normal.element.get_or_add_rPr()
    r_fonts = r_pr.find(qn("w:rFonts"))
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.append(r_fonts)
    for attribute in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        r_fonts.set(qn(attribute), FONT)

    bullet = document.styles["List Bullet"]
    bullet.font.name = FONT
    bullet.font.size = Pt(10.5)
    bullet.font.color.rgb = rgb(INK)
    fmt = bullet.paragraph_format
    fmt.left_indent = Cm(0.45)
    fmt.first_line_indent = Cm(-0.3)
    fmt.space_before = Pt(0)
    fmt.space_after = Pt(1)
    fmt.line_spacing = 1.0

    props = document.core_properties
    props.author = "Nosheene Mohammad"
    props.title = "CV — Nosheene Mohammad — Développeuse web full stack"
    props.subject = "Curriculum vitae"
    props.category = "CV"
    props.language = "fr-FR"


def add_section_title(document: Document, text: str) -> None:
    paragraph = document.add_paragraph()
    set_paragraph_spacing(paragraph, before=9, after=2, line=230)
    add_text(paragraph, text.upper(), 11, bold=True, color=ACCENT)
    add_bottom_border(paragraph, ACCENT, "8")
    paragraph.paragraph_format.keep_with_next = True


def add_role(document: Document, title: str, organization: str, period: str) -> None:
    paragraph = document.add_paragraph()
    set_paragraph_spacing(paragraph, before=6, after=0, line=230)
    paragraph.paragraph_format.tab_stops.add_tab_stop(Cm(17.8), WD_TAB_ALIGNMENT.RIGHT)
    paragraph.paragraph_format.keep_with_next = True
    add_text(paragraph, title, 11, bold=True)
    add_text(paragraph, "\t" + period, 11, color=ACCENT)

    place = document.add_paragraph()
    set_paragraph_spacing(place, before=0, after=1, line=210)
    place.paragraph_format.keep_with_next = True
    add_text(place, organization, 10.5, color=MUTED)


def add_bullet(document: Document, text: str) -> None:
    paragraph = document.add_paragraph(style="List Bullet")
    set_paragraph_spacing(paragraph, before=0, after=1, line=226)
    paragraph.clear()
    add_text(paragraph, text, 10.5)


def add_skill_line(document: Document, label: str, content: str) -> None:
    paragraph = document.add_paragraph()
    set_paragraph_spacing(paragraph, before=1, after=1, line=226)
    add_text(paragraph, label, 10.5, bold=True)
    add_text(paragraph, content, 10.5)


def add_entry_line(document: Document, primary: str, secondary: str, period: str) -> None:
    paragraph = document.add_paragraph()
    set_paragraph_spacing(paragraph, before=1, after=1, line=226)
    paragraph.paragraph_format.tab_stops.add_tab_stop(Cm(17.8), WD_TAB_ALIGNMENT.RIGHT)
    add_text(paragraph, primary, 10.5, bold=True)
    if secondary:
        add_text(paragraph, "  —  " + secondary, 10.5, color=MUTED)
    add_text(paragraph, "\t" + period, 10.5, color=ACCENT)


def build() -> Document:
    document = Document()
    configure_document(document)

    name = document.add_paragraph()
    set_paragraph_spacing(name, before=0, after=0, line=360)
    name.paragraph_format.keep_with_next = True
    add_left_border(name)
    add_text(name, "Nosheene Mohammad", 22, bold=True, color=INK)

    role = document.add_paragraph()
    set_paragraph_spacing(role, before=0, after=1, line=240)
    role.paragraph_format.keep_with_next = True
    add_left_border(role)
    add_text(role, "Développeuse web full stack", 12, bold=True, color=ACCENT)

    contact = document.add_paragraph()
    set_paragraph_spacing(contact, before=1, after=0, line=210)
    contact.paragraph_format.keep_with_next = True
    add_left_border(contact)
    add_text(contact, "Ozoir-la-Ferrière  ·  06 84 47 71 19  ·  ", 10, color=MUTED)
    add_hyperlink(contact, "mohammadnosheene@gmail.com", "mailto:mohammadnosheene@gmail.com", 10)

    links = document.add_paragraph()
    set_paragraph_spacing(links, before=0, after=2, line=210)
    add_left_border(links)
    add_hyperlink(
        links,
        "linkedin.com/in/nosheene-mohammad-498980160",
        "https://www.linkedin.com/in/nosheene-mohammad-498980160",
        10,
    )
    add_text(links, "  ·  RQTH", 10, color=MUTED)
    add_bottom_border(links, RULE, "8")

    add_section_title(document, "Profil")
    profile = document.add_paragraph()
    set_paragraph_spacing(profile, before=2, after=1, line=226)
    add_text(
        profile,
        "Développeuse web full stack, à l’aise avec le cadrage, l’intégration front-end et la qualité logicielle. "
        "De la maquette au déploiement\u00a0: JavaScript, échanges d’API, SQL, tests fonctionnels, "
        "tests d’API avec Postman, recette et analyse des anomalies. Capacité d’analyse et de synthèse, "
        "créative, rigoureuse et organisée.",
        10.5,
    )

    add_section_title(document, "Compétences")
    add_skill_line(
        document,
        "Développement web\u00a0: ",
        "JavaScript, intégration front-end, interactivité, Fetch, JSON, API REST, SQL, "
        "bases de données relationnelles, architecture web, Git, déploiement.",
    )
    add_skill_line(
        document,
        "Qualité logicielle\u00a0: ",
        "tests fonctionnels, recette, contrôle qualité, débogage, analyse d’anomalies, "
        "tests d’API (Postman), RGAA, normes de sécurité.",
    )
    add_skill_line(
        document,
        "Conception\u00a0: ",
        "cadrage et suivi de projet, maquettage, prototypage, UX design, Figma, Draw.io.",
    )
    add_skill_line(
        document,
        "Pédagogie digitale\u00a0: ",
        "conception e-learning, Articulate Storyline 360, Rise 360, LMS, reporting.",
    )

    add_section_title(document, "Expérience professionnelle")
    add_role(document, "Intégratrice e-learning", "Formalearning", "En cours")
    add_bullet(document, "Intégrer des modules e-learning dans Storyline 360 et Rise 360.")
    add_bullet(document, "Assurer le support client LMS\u00a0: recette des parcours et résolution des bugs.")

    add_role(
        document,
        "Conceptrice pédagogique digitale",
        "Sonepar France Interservices",
        "2019 – 2023",
    )
    add_bullet(
        document,
        "Concevoir et intégrer des modules e-learning (Storyline, LCMS Talentsoft, Rise 360), puis réaliser la recette.",
    )
    add_bullet(
        document,
        "Administrer le LMS Talentsoft et LinkedIn Learning pour 5\u202f600 collaborateurs\u00a0: "
        "inventaire, mise à jour et publication des contenus de formation.",
    )
    add_bullet(document, "Produire les reportings mensuels, annuels et spécifiques.")
    add_bullet(
        document,
        "Rédiger et publier les communications internes et le catalogue de formation e-learning.",
    )

    add_role(document, "Conceptrice pédagogique digitale", "Tuto’s Me Pro", "2019")
    add_bullet(
        document,
        "Rédiger la note de cadrage d’une formation multimodale «\u00a0Formateur professionnel\u00a0».",
    )
    add_bullet(document, "Concevoir le storyboard de la formation.")

    add_role(
        document,
        "Formatrice professionnelle pour adultes",
        "AFPS, Le Panse Academy, Le Panse Formation, GDM Formations, AFEC",
        "2016 – 2018",
    )
    add_bullet(document, "Concevoir et animer des formations en présentiel, et mettre en place un dispositif tutoral.")
    add_bullet(document, "Gérer les plateformes interministérielles Foromes et Chronos (BPJEPS AGFF/AF).")

    add_section_title(document, "Formation")
    add_entry_line(document, "Développeur web full stack", "Studi", "2026")
    add_entry_line(document, "Jeux, sound design, droits d’auteur et JavaScript", "", "2023 – 2025")
    add_entry_line(document, "UX design", "", "2021")
    add_entry_line(document, "Concepteur de contenus digitaux de formation", "Evocime", "2019")
    add_entry_line(document, "Titre de formateur professionnel pour adultes", "AFPA", "2017")

    add_section_title(document, "Outils")
    tools = document.add_paragraph()
    set_paragraph_spacing(tools, before=2, after=1, line=226)
    add_text(
        tools,
        "JavaScript, Postman, Git, SQL, Fetch, Docker, Figma, Draw.io, Articulate Storyline 360, "
        "Rise 360, LMS Cegid Talentsoft, LCMS Talentsoft, Microsoft Office, Microsoft Forms, Sway, Loop. "
        "Environnements\u00a0: PC et Mac.",
        10.5,
    )

    add_section_title(document, "Langues et centre d’intérêt")
    extra = document.add_paragraph()
    set_paragraph_spacing(extra, before=2, after=0, line=226)
    add_text(extra, "Anglais\u00a0: niveau B2.", 10.5)
    interest = document.add_paragraph()
    set_paragraph_spacing(interest, before=0, after=0, line=226)
    add_text(interest, "Centre d’intérêt\u00a0: Tennis.", 10.5)

    document.save(OUTPUT)
    return document


if __name__ == "__main__":
    build()
    print(OUTPUT)
